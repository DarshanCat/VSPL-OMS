import openpyxl
import os
import pandas as pd

excel_path = r"C:\Users\Darshan\.gemini\antigravity\brain\9283038f-a0e0-4fca-b665-b36c5e21dba1\.user_uploaded\media_1791441452333.xlsx"

wb = openpyxl.load_workbook(excel_path, data_only=True)
print("Sheet names:", wb.sheetnames)

for sheetname in wb.sheetnames:
    ws = wb[sheetname]
    print(f"\n--- SHEET: {sheetname} ---")
    print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")
    
    # Read headers
    headers = [cell.value for cell in ws[1]]
    print(f"Headers ({len(headers)}):", headers)
    
    # Read first 5 data rows
    print("\nFirst 5 rows:")
    for r in range(2, min(7, ws.max_row + 1)):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        print(f"Row {r}:", row_vals[:10])

# Also let's load with pandas to get total non-empty rows, etc.
df = pd.read_excel(excel_path, sheet_name=0)
print(f"\nPandas shape: {df.shape}")
print(f"Columns: {list(df.columns)}")
print(f"Non-null counts:\n{df.count()}")
