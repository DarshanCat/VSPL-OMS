import openpyxl
import subprocess
import json
import uuid
from datetime import datetime, date
from uuid import UUID
from urllib.parse import urlparse, unquote
from google.cloud.sql.connector import Connector, IPTypes
from google.oauth2.credentials import Credentials
import pg8000.dbapi
from collections import defaultdict

class CustomEncoder(json.JSONEncoder):
    def default(self, obj):
        if isinstance(obj, UUID):
            return str(obj)
        if isinstance(obj, (datetime, date)):
            return obj.isoformat()
        return super().default(obj)

excel_path = r"C:\Users\Darshan\.gemini\antigravity\brain\9283038f-a0e0-4fca-b665-b36c5e21dba1\.user_uploaded\media_1791441452333.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Sheet1']

rows = []
for r in range(2, ws.max_row + 1):
    row_data = [ws.cell(row=r, column=c).value for c in range(1, 23)]
    if not any(row_data):
        continue
    rows.append(row_data)

headers = [
    'wo', 'ts', 'oar', 'cust', 'cust_code', 'uniq_prod_id',
    'part_no', 'internal_part_code', 'internal_id', 'region',
    'grade', 'cust_grade', 'order_type', 'order_cat',
    'po_num', 'del_date', 'po_rate', 'po_qty', 'wo_qty',
    'remarks', 'wo_type', 'status'
]

source_data = [dict(zip(headers, r)) for r in rows]
print(f"Loaded {len(source_data)} source rows from Excel.")

# Connect to Cloud SQL DB
print("Connecting to Cloud SQL (vspl-oms-db)...")
token = subprocess.check_output('gcloud.cmd auth print-access-token', shell=True).decode().strip()
credentials = Credentials(token)

secret = subprocess.check_output('gcloud.cmd secrets versions access latest --secret=VSPL_DATABASE_URL --project=vspl-oms-510015', shell=True).decode().strip()
parsed = urlparse(secret)
username = unquote(parsed.username) if parsed.username else 'vspl_user'
password = unquote(parsed.password) if parsed.password else ''
dbname = parsed.path.lstrip('/')

connector = Connector(credentials=credentials)
conn = connector.connect(
    'vspl-oms-510015:asia-south1:vspl-oms-db',
    'pg8000',
    user=username,
    password=password,
    db=dbname,
    ip_type=IPTypes.PUBLIC
)
conn.autocommit = False
cur = conn.cursor()

print("==================================================")
print("1. PRE-IMPORT VERIFICATION & SNAPSHOT")
print("==================================================")

# Check baseline tables
cur.execute("SELECT id, customer_code, name FROM customers;")
db_customers = {r[1]: (r[0], r[2]) for r in cur.fetchall()}

cur.execute("SELECT id, part_number, description FROM parts;")
db_parts = {r[1]: (r[0], r[2]) for r in cur.fetchall()}

cur.execute("SELECT COUNT(*) FROM po_master;")
baseline_po_master = cur.fetchone()[0]
print(f"Baseline po_master count: {baseline_po_master}")

cur.execute("SELECT COUNT(*) FROM po_lines;")
baseline_po_lines = cur.fetchone()[0]
print(f"Baseline po_lines count: {baseline_po_lines}")

cur.execute("SELECT COUNT(*) FROM orders;")
baseline_orders = cur.fetchone()[0]
print(f"Baseline orders count: {baseline_orders}")

cur.execute("SELECT COUNT(*) FROM work_orders;")
baseline_wos = cur.fetchone()[0]
print(f"Baseline work_orders count: {baseline_wos}")

protected_baselines = {
    'customers': 119,
    'parts': 5506,
    'customer_part_mappings': 5170,
    'customer_part_cross_references': 1891,
    'cc_materials': 15,
    'cc_inwards': 15,
    'cc_stock_units': 273,
    'cc_stock_ledger': 273,
    'users': 27,
    'audit_logs': 332
}

for t, exp in protected_baselines.items():
    cur.execute(f"SELECT COUNT(*) FROM {t};")
    cnt = cur.fetchone()[0]
    assert cnt == exp, f"Protected table {t} count mismatch: got {cnt}, expected {exp}"
    print(f"  [OK] {t}: {cnt}")

print("\n==================================================")
print("2. BUILDING IMPORT DATA STRUCTURES")
print("==================================================")

# Group by (customer_code, po_num) -> 20 PO headers
po_headers_map = {} # (cust_code, po_num) -> po_dict
po_lines_map = {}   # (cust_code, po_num, internal_id) -> po_line_dict
orders_map = {}     # oar_number -> order_dict
wos_list = []       # list of wo_dicts

for d in source_data:
    cust_code = str(d['cust_code']).strip()
    po_num = str(d['po_num']).strip()
    internal_id = str(d['internal_id']).strip()
    oar_num = str(d['oar']).strip()
    wo_num = str(d['wo']).strip()
    
    cust_id, cust_name = db_customers[cust_code]
    part_id, part_desc = db_parts[internal_id]
    
    po_key = (cust_code, po_num)
    if po_key not in po_headers_map:
        po_headers_map[po_key] = {
            "id": str(uuid.uuid4()),
            "po_number": po_num,
            "customer_id": cust_id,
            "customer_code": cust_code,
            "customer_name": cust_name,
            "po_date": d['del_date'] if isinstance(d['del_date'], (date, datetime)) else None,
            "validity_date": d['del_date'] if isinstance(d['del_date'], (date, datetime)) else None,
            "status": "OPEN",
            "created_at": d['ts'] if isinstance(d['ts'], (date, datetime)) else datetime.utcnow()
        }
    
    po_line_key = (cust_code, po_num, internal_id)
    if po_line_key not in po_lines_map:
        po_lines_map[po_line_key] = {
            "id": str(uuid.uuid4()),
            "po_id": po_headers_map[po_key]["id"],
            "part_id": part_id,
            "internal_id": internal_id,
            "po_qty": int(d['po_qty'] or 0),
            "required_date": d['del_date'] if isinstance(d['del_date'], (date, datetime)) else None,
            "created_at": d['ts'] if isinstance(d['ts'], (date, datetime)) else datetime.utcnow()
        }
        
    if oar_num not in orders_map:
        orders_map[oar_num] = {
            "id": str(uuid.uuid4()),
            "oar_number": oar_num,
            "customer_id": cust_id,
            "part_id": part_id,
            "customer_po": po_num,
            "po_qty": int(d['po_qty'] or 0),
            "max_batch_size": 100,
            "delivery_date": d['del_date'] if isinstance(d['del_date'], (date, datetime)) else None,
            "order_type": d['order_type'] or "PO",
            "status": "CONFIRMED",
            "source_type": "po",
            "po_line_id": po_lines_map[po_line_key]["id"],
            "order_classification": d['order_cat'] or "Serial",
            "created_at": d['ts'] if isinstance(d['ts'], (date, datetime)) else datetime.utcnow()
        }
    
    wo_dict = {
        "id": str(uuid.uuid4()),
        "wo_number": wo_num,
        "order_id": orders_map[oar_num]["id"],
        "physical_wo_qty": int(d['wo_qty'] or 0),
        "current_stage": "F1",
        "projected_final_good": int(d['wo_qty'] or 0),
        "shortfall": "No",
        "status": str(d['status']).strip() if d['status'] and str(d['status']).strip() else "OPEN",
        "source_wo_id": None,
        "is_replacement": (str(d['wo_type']).strip().lower() == "conversion"),
        "casting_process": d['grade'],
        "remarks": d['remarks'] or "",
        "created_at": d['ts'] if isinstance(d['ts'], (date, datetime)) else datetime.utcnow()
    }
    wos_list.append(wo_dict)

print(f"Constructed: {len(po_headers_map)} PO Headers, {len(po_lines_map)} PO Lines, {len(orders_map)} Orders (OARs), {len(wos_list)} Work Orders")
assert len(po_headers_map) == 20, f"Expected 20 PO headers, got {len(po_headers_map)}"
assert len(po_lines_map) == 41, f"Expected 41 PO lines, got {len(po_lines_map)}"
assert len(orders_map) == 41, f"Expected 41 Orders, got {len(orders_map)}"
assert len(wos_list) == 59, f"Expected 59 Work Orders, got {len(wos_list)}"

# Verify no target records exist in DB prior to insertion
print("\n--- TARGET CONFLICT PRE-CHECKS ---")
target_oars = list(orders_map.keys())
cur.execute("SELECT COUNT(*) FROM orders WHERE oar_number = ANY(%s);", (target_oars,))
existing_oar_count = cur.fetchone()[0]
assert existing_oar_count == 0, f"Found {existing_oar_count} existing target OARs in DB!"
print(f"  [OK] Zero target OARs pre-exist in DB.")

target_wos = [w['wo_number'] for w in wos_list]
cur.execute("SELECT COUNT(*) FROM work_orders WHERE wo_number = ANY(%s);", (target_wos,))
existing_wo_count = cur.fetchone()[0]
assert existing_wo_count == 0, f"Found {existing_wo_count} existing target WOs in DB!"
print(f"  [OK] Zero target WOs pre-exist in DB.")

# Save pre-import snapshot
pre_snapshot = {
    "timestamp": datetime.utcnow().isoformat() + "Z",
    "baseline_counts": {
        "po_master": baseline_po_master,
        "po_lines": baseline_po_lines,
        "orders": baseline_orders,
        "work_orders": baseline_wos
    },
    "target_counts": {
        "po_master": len(po_headers_map),
        "po_lines": len(po_lines_map),
        "orders": len(orders_map),
        "work_orders": len(wos_list)
    },
    "protected_tables": protected_baselines
}
with open(r"c:\VSPL\VSPL smart\scratch\pre_import_snapshot_59wos.json", "w", encoding="utf-8") as f:
    json.dump(pre_snapshot, f, indent=2)
print("Pre-import snapshot written to scratch/pre_import_snapshot_59wos.json")

print("\n==================================================")
print("3. EXECUTING ATOMIC TRANSACTION")
print("==================================================")

try:
    # Update index on po_master to support multi-customer PO numbering
    cur.execute("DROP INDEX IF EXISTS ix_po_master_po_number;")
    cur.execute("CREATE UNIQUE INDEX IF NOT EXISTS ix_po_master_customer_po ON public.po_master USING btree (customer_id, po_number);")
    print("Updated po_master index to UNIQUE (customer_id, po_number)")

    # 1. Insert 20 po_master rows
    for p in po_headers_map.values():
        cur.execute("""
        INSERT INTO po_master (id, po_number, customer_id, po_date, validity_date, status, created_by_name, created_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s);
        """, (p['id'], p['po_number'], p['customer_id'], p['po_date'], p['validity_date'], p['status'], 'system:import', p['created_at']))
    print(f"Inserted {len(po_headers_map)} po_master records.")

    # 2. Insert 41 po_lines rows
    for l in po_lines_map.values():
        cur.execute("""
        INSERT INTO po_lines (id, po_id, part_id, po_qty, required_date, created_at)
        VALUES (%s, %s, %s, %s, %s, %s);
        """, (l['id'], l['po_id'], l['part_id'], l['po_qty'], l['required_date'], l['created_at']))
    print(f"Inserted {len(po_lines_map)} po_lines records.")

    # 3. Insert 41 orders rows
    for o in orders_map.values():
        cur.execute("""
        INSERT INTO orders (id, oar_number, customer_id, part_id, customer_po, po_qty, max_batch_size, delivery_date, order_type, status, source_type, po_line_id, order_classification, created_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
        """, (o['id'], o['oar_number'], o['customer_id'], o['part_id'], o['customer_po'], o['po_qty'], o['max_batch_size'], o['delivery_date'], o['order_type'], o['status'], o['source_type'], o['po_line_id'], o['order_classification'], o['created_at']))
    print(f"Inserted {len(orders_map)} orders records.")

    # 4. Insert 59 work_orders rows
    for w in wos_list:
        cur.execute("""
        INSERT INTO work_orders (id, wo_number, order_id, physical_wo_qty, current_stage, projected_final_good, shortfall, status, source_wo_id, is_replacement, casting_process, manufacturing_remarks, created_at)
        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s);
        """, (w['id'], w['wo_number'], w['order_id'], w['physical_wo_qty'], w['current_stage'], w['projected_final_good'], w['shortfall'], w['status'], w['source_wo_id'], w['is_replacement'], w['casting_process'], w['remarks'], w['created_at']))
    print(f"Inserted {len(wos_list)} work_orders records.")

    # In-transaction assertions
    cur.execute("SELECT COUNT(*) FROM po_master;")
    assert cur.fetchone()[0] == baseline_po_master + 20, "po_master count mismatch inside txn!"

    cur.execute("SELECT COUNT(*) FROM po_lines;")
    assert cur.fetchone()[0] == baseline_po_lines + 41, "po_lines count mismatch inside txn!"

    cur.execute("SELECT COUNT(*) FROM orders;")
    assert cur.fetchone()[0] == baseline_orders + 41, "orders count mismatch inside txn!"

    cur.execute("SELECT COUNT(*) FROM work_orders;")
    assert cur.fetchone()[0] == baseline_wos + 59, "work_orders count mismatch inside txn!"

    for t, exp in protected_baselines.items():
        cur.execute(f"SELECT COUNT(*) FROM {t};")
        cnt = cur.fetchone()[0]
        assert cnt == exp, f"Protected table {t} mutated inside txn: got {cnt}, expected {exp}"

    # Commit
    conn.commit()
    print("\n>>> TRANSACTION COMMITTED SUCCESSFULLY <<<")

except Exception as e:
    conn.rollback()
    print(f"\n!!! TRANSACTION FAILED & ROLLED BACK: {e} !!!")
    raise

print("\n==================================================")
print("4. POST-COMMIT RECONCILIATION & VALIDATION")
print("==================================================")

# Query back all WOs with complete joins
cur.execute("""
SELECT 
    w.wo_number,
    w.physical_wo_qty,
    w.status AS wo_status,
    w.is_replacement,
    o.oar_number,
    o.customer_po,
    o.po_qty,
    o.delivery_date,
    c.customer_code,
    c.name AS customer_name,
    p.part_number,
    p.description AS part_desc
FROM work_orders w
JOIN orders o ON w.order_id = o.id
JOIN customers c ON o.customer_id = c.id
JOIN parts p ON o.part_id = p.id
WHERE w.wo_number = ANY(%s)
ORDER BY w.wo_number ASC;
""", (target_wos,))
db_wos_joined = cur.fetchall()
print(f"Total joined WOs retrieved from database: {len(db_wos_joined)}")
assert len(db_wos_joined) == 59, f"Expected 59 joined WOs, got {len(db_wos_joined)}"

# Reconcile against source_data
source_by_wo = {d['wo']: d for d in source_data}
reconciliation_mismatches = []

for r in db_wos_joined:
    wo_num = r[0]
    src = source_by_wo.get(wo_num)
    if not src:
        reconciliation_mismatches.append(f"WO {wo_num} in DB but not in source!")
        continue
    
    # Check WO Qty
    if r[1] != int(src['wo_qty']):
        reconciliation_mismatches.append(f"WO {wo_num} qty mismatch: DB={r[1]}, Source={src['wo_qty']}")
    # Check OAR
    if str(r[4]).strip() != str(src['oar']).strip():
        reconciliation_mismatches.append(f"WO {wo_num} OAR mismatch: DB={r[4]}, Source={src['oar']}")
    # Check Customer PO (string normalized)
    if str(r[5]).strip() != str(src['po_num']).strip():
        reconciliation_mismatches.append(f"WO {wo_num} PO mismatch: DB={r[5]}, Source={src['po_num']}")
    # Check Customer Code
    if str(r[8]).strip() != str(src['cust_code']).strip():
        reconciliation_mismatches.append(f"WO {wo_num} Customer Code mismatch: DB={r[8]}, Source={src['cust_code']}")
    # Check Part Number
    if str(r[10]).strip() != str(src['internal_id']).strip():
        reconciliation_mismatches.append(f"WO {wo_num} Part mismatch: DB={r[10]}, Source={src['internal_id']}")

print(f"Source-to-DB reconciliation errors: {len(reconciliation_mismatches)}")
if reconciliation_mismatches:
    for err in reconciliation_mismatches:
        print("  ERROR:", err)
    raise AssertionError("Reconciliation failed!")
else:
    print("100% PERFECT 59/59 WORK ORDER RECONCILIATION!")

# Check Quantity Exceptions Preservation
print("\n--- QUANTITY EXCEPTIONS CHECK ---")
# 1. WO-17476
cur.execute("""
SELECT w.wo_number, w.physical_wo_qty, o.po_qty, p.part_number, o.customer_po
FROM work_orders w
JOIN orders o ON w.order_id = o.id
JOIN parts p ON o.part_id = p.id
WHERE w.wo_number = 'WO-17476';
""")
r1 = cur.fetchone()
print(f"Exception 1 (WO-17476): WO={r1[0]}, Part={r1[3]}, PO={r1[4]}, PO Qty={r1[2]}, WO Qty={r1[1]} (Preserved exact: PO Qty 2, WO Qty 10)")
assert r1[1] == 10 and r1[2] == 2

# 2. AUM3 / PO 9803536
cur.execute("""
SELECT w.wo_number, w.physical_wo_qty, o.po_qty, p.part_number, o.customer_po
FROM work_orders w
JOIN orders o ON w.order_id = o.id
JOIN parts p ON o.part_id = p.id
WHERE o.customer_po = '9803536' AND p.part_number = 'AUM3'
ORDER BY w.wo_number ASC;
""")
r2 = cur.fetchall()
print(f"Exception 2 (AUM3 / 9803536):")
tot_aum3_wo = sum(row[1] for row in r2)
for row in r2:
    print(f"  WO={row[0]}, Part={row[3]}, PO={row[4]}, PO Qty={row[2]}, WO Qty={row[1]}")
print(f"  Total WO Qty for AUM3 = {tot_aum3_wo} against PO Qty {r2[0][2]} (Preserved exact: PO Qty 2, Total WO Qty 4)")
assert tot_aum3_wo == 4 and r2[0][2] == 2

# Check Final Protected Counts
print("\nFinal Protected Masters Check:")
for t, exp in protected_baselines.items():
    cur.execute(f"SELECT COUNT(*) FROM {t};")
    cnt = cur.fetchone()[0]
    print(f"  {t}: {cnt} (expected {exp})")
    assert cnt == exp

cur.close()
conn.close()
connector.close()
print("\nALL POST-IMPORT VERIFICATIONS COMPLETED SUCCESSFULLY!")
