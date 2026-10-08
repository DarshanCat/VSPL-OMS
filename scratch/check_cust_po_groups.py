import openpyxl
from collections import defaultdict

excel_path = r"C:\Users\Darshan\.gemini\antigravity\brain\9283038f-a0e0-4fca-b665-b36c5e21dba1\.user_uploaded\media_1791441452333.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Sheet1']

rows = []
for r in range(2, ws.max_row + 1):
    row_data = [ws.cell(row=r, column=c).value for c in range(1, 23)]
    if not any(row_data):
        continue
    rows.append(row_data)

headers = [
    'wo', 'ts', 'oar', 'cust', 'cust_code', 'uniq_prod_id',
    'part_no', 'internal_part_code', 'internal_id', 'region',
    'grade', 'cust_grade', 'order_type', 'order_cat',
    'po_num', 'del_date', 'po_rate', 'po_qty', 'wo_qty',
    'remarks', 'wo_type', 'status'
]

data = [dict(zip(headers, r)) for r in rows]

# Group by (Customer Code, PO Number)
cust_po_groups = defaultdict(list)
for d in data:
    cust_code = str(d['cust_code']).strip()
    po_num = str(d['po_num']).strip()
    cust_po_groups[(cust_code, po_num)].append(d)

print(f"Total Unique (Customer Code, PO Number) pairs: {len(cust_po_groups)}")
for (c_code, po_num), recs in sorted(cust_po_groups.items()):
    cust_name = recs[0]['cust']
    oars = sorted(list(set(r['oar'] for r in recs)))
    wos = [r['wo'] for r in recs]
    tot_wo_qty = sum(float(r['wo_qty'] or 0) for r in recs)
    print(f"\nCustomer: {c_code} ({cust_name}) | PO: '{po_num}'")
    print(f"  OARs ({len(oars)}): {oars}")
    print(f"  WOs ({len(wos)}): {wos} | Total WO Qty: {tot_wo_qty}")
