import openpyxl
import pandas as pd
from collections import defaultdict
import json

excel_path = r"C:\Users\Darshan\.gemini\antigravity\brain\9283038f-a0e0-4fca-b665-b36c5e21dba1\.user_uploaded\media_1791441452333.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Sheet1']

rows = []
for r in range(2, ws.max_row + 1):
    row_data = [ws.cell(row=r, column=c).value for c in range(1, 23)]
    if not any(row_data):
        continue
    rows.append(row_data)

headers = [
    'wo', 'ts', 'oar', 'cust', 'cust_code', 'uniq_prod_id',
    'part_no', 'internal_part_code', 'internal_id', 'region',
    'grade', 'cust_grade', 'order_type', 'order_cat',
    'po_num', 'del_date', 'po_rate', 'po_qty', 'wo_qty',
    'remarks', 'wo_type', 'status'
]

data = [dict(zip(headers, r)) for r in rows]

# Build hierarchical structure: PO -> OAR -> WO
hierarchy = defaultdict(lambda: {
    "customer": None,
    "customer_code": None,
    "oars": defaultdict(lambda: {
        "part_no": None,
        "internal_id": None,
        "wos": []
    })
})

for d in data:
    po = str(d['po_num']).strip()
    oar = str(d['oar']).strip()
    wo = str(d['wo']).strip()
    
    hierarchy[po]["customer"] = d['cust']
    hierarchy[po]["customer_code"] = d['cust_code']
    
    hierarchy[po]["oars"][oar]["part_no"] = d['part_no']
    hierarchy[po]["oars"][oar]["internal_id"] = d['internal_id']
    hierarchy[po]["oars"][oar]["wos"].append(d)

import sys
sys.stdout.reconfigure(encoding='utf-8')

print("Total POs in hierarchy:", len(hierarchy))
for po, pinfo in hierarchy.items():
    print(f"\nPO: '{po}' ({pinfo['customer_code']} - {pinfo['customer']}) | OARs: {len(pinfo['oars'])}")
    for oar, oinfo in pinfo['oars'].items():
        tot_wo_qty = sum(float(w['wo_qty'] or 0) for w in oinfo['wos'])
        po_qty = oinfo['wos'][0]['po_qty']
        print(f"  +-- OAR: '{oar}' | Part: {oinfo['internal_id']} ({oinfo['part_no']}) | PO Qty: {po_qty} | WOs ({len(oinfo['wos'])}): Total WO Qty = {tot_wo_qty}")
        for w in oinfo['wos']:
            print(f"       +-- WO: {w['wo']} | Qty: {w['wo_qty']} | Type: {w['wo_type']} | Status: {w['status']} | Del Date: {str(w['del_date'])[:10]}")
