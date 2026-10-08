import openpyxl
import subprocess
from urllib.parse import urlparse, unquote
from google.cloud.sql.connector import Connector, IPTypes
from google.oauth2.credentials import Credentials
import pg8000.dbapi

excel_path = r"C:\Users\Darshan\.gemini\antigravity\brain\9283038f-a0e0-4fca-b665-b36c5e21dba1\.user_uploaded\media_1791441452333.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Sheet1']

rows = []
for r in range(2, ws.max_row + 1):
    row_data = [ws.cell(row=r, column=c).value for c in range(1, 23)]
    if not any(row_data):
        continue
    rows.append(row_data)

token = subprocess.check_output('gcloud.cmd auth print-access-token', shell=True).decode().strip()
credentials = Credentials(token)

secret = subprocess.check_output('gcloud.cmd secrets versions access latest --secret=VSPL_DATABASE_URL --project=vspl-oms-510015', shell=True).decode().strip()
parsed = urlparse(secret)
username = unquote(parsed.username) if parsed.username else 'vspl_user'
password = unquote(parsed.password) if parsed.password else ''
dbname = parsed.path.lstrip('/')

connector = Connector(credentials=credentials)
conn = connector.connect(
    'vspl-oms-510015:asia-south1:vspl-oms-db',
    'pg8000',
    user=username,
    password=password,
    db=dbname,
    ip_type=IPTypes.PUBLIC
)
cur = conn.cursor()

print("--- PART MATCHING DETAILED AUDIT ---")
part_results = []

for idx, r in enumerate(rows, 1):
    wo = str(r[0]).strip()
    cust_code = str(r[4]).strip()
    part_no = str(r[6]).strip() if r[6] else ""
    internal_part_code = str(r[7]).strip() if r[7] else ""
    internal_id = str(r[8]).strip() if r[8] else ""
    
    # 1. Exact match on internal_id == parts.part_number
    cur.execute("SELECT id, part_number, description FROM parts WHERE part_number = %s;", (internal_id,))
    p1 = cur.fetchall()
    
    # 2. Exact match on part_no == parts.part_number
    cur.execute("SELECT id, part_number, description FROM parts WHERE part_number = %s;", (part_no,))
    p2 = cur.fetchall()
    
    # 3. Check customer_part_mappings
    cur.execute("""
    SELECT p.id, p.part_number, m.customer_part_number
    FROM customer_part_mappings m
    JOIN parts p ON m.part_id = p.id
    JOIN customers c ON m.customer_id = c.id
    WHERE c.customer_code = %s AND (m.customer_part_number = %s OR p.part_number = %s);
    """, (cust_code, part_no, internal_id))
    p3 = cur.fetchall()
    
    # 4. Check customer_part_cross_references
    cur.execute("""
    SELECT p.id, p.part_number, cr.customer_part_no
    FROM customer_part_cross_references cr
    JOIN parts p ON cr.part_id = p.id
    JOIN customers c ON cr.customer_id = c.id
    WHERE c.customer_code = %s AND (cr.customer_part_no = %s OR p.part_number = %s);
    """, (cust_code, part_no, internal_id))
    p4 = cur.fetchall()
    
    part_results.append({
        "wo": wo,
        "cust_code": cust_code,
        "part_no": part_no,
        "internal_part_code": internal_part_code,
        "internal_id": internal_id,
        "match_internal_id": p1,
        "match_part_no": p2,
        "match_mapping": p3,
        "match_cross_ref": p4
    })

# Summarize matching
all_internal_id_matched = all(len(p['match_internal_id']) > 0 for p in part_results)
print(f"All 59 rows match parts.part_number == internal_id: {all_internal_id_matched}")

unique_parts = set()
for p in part_results:
    unique_parts.add((p['cust_code'], p['internal_id'], p['part_no']))

print(f"\nUnique Parts across the 59 WOs: {len(unique_parts)}")
for cust_code, int_id, cust_part in sorted(unique_parts):
    # Get DB part info
    cur.execute("SELECT id, part_number, description FROM parts WHERE part_number = %s;", (int_id,))
    db_p = cur.fetchone()
    if db_p:
        print(f"Customer: {cust_code:4} | Internal ID: {int_id:8} | Customer Part: {cust_part:35} | DB Part: '{db_p[1]}' (UUID: {db_p[0]})")
    else:
        print(f"Customer: {cust_code:4} | Internal ID: {int_id:8} | Customer Part: {cust_part:35} | NOT FOUND IN DB!")

cur.close()
conn.close()
connector.close()
