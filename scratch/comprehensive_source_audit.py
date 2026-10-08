import openpyxl
import os
import pandas as pd
import subprocess
from urllib.parse import urlparse, unquote
from google.cloud.sql.connector import Connector, IPTypes
from google.oauth2.credentials import Credentials
import pg8000.dbapi
from collections import defaultdict

excel_path = r"C:\Users\Darshan\.gemini\antigravity\brain\9283038f-a0e0-4fca-b665-b36c5e21dba1\.user_uploaded\media_1791441452333.xlsx"

wb = openpyxl.load_workbook(excel_path, data_only=True)
sheetnames = wb.sheetnames
ws = wb[sheetnames[0]]

headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

rows = []
for r in range(2, ws.max_row + 1):
    row_data = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if not any(row_data):
        continue
    rows.append(row_data)

print(f"Total extracted rows: {len(rows)}")

# Connect to DB
token = subprocess.check_output('gcloud.cmd auth print-access-token', shell=True).decode().strip()
credentials = Credentials(token)

secret = subprocess.check_output('gcloud.cmd secrets versions access latest --secret=VSPL_DATABASE_URL --project=vspl-oms-510015', shell=True).decode().strip()
parsed = urlparse(secret)
username = unquote(parsed.username) if parsed.username else 'vspl_user'
password = unquote(parsed.password) if parsed.password else ''
dbname = parsed.path.lstrip('/')

connector = Connector(credentials=credentials)
conn = connector.connect(
    'vspl-oms-510015:asia-south1:vspl-oms-db',
    'pg8000',
    user=username,
    password=password,
    db=dbname,
    ip_type=IPTypes.PUBLIC
)
cur = conn.cursor()

# Check table columns
for t in ['customer_part_cross_references', 'customer_part_mappings', 'parts']:
    cur.execute("""
    SELECT column_name, data_type 
    FROM information_schema.columns 
    WHERE table_name = %s 
    ORDER BY ordinal_position;
    """, (t,))
    print(t, 'columns:', [r[0] for r in cur.fetchall()])

# 1. Fetch DB customers
cur.execute("SELECT id, customer_code, name FROM customers;")
db_customers = {r[1]: (r[0], r[2]) for r in cur.fetchall()}

# 2. Fetch DB parts
cur.execute("SELECT id, part_number, description FROM parts;")
db_parts = {r[1]: (r[0], r[2]) for r in cur.fetchall()}

# 3. Fetch DB customer_part_mappings
cur.execute("""
SELECT m.id, c.customer_code, p.part_number, m.customer_part_number, m.part_id, m.customer_id
FROM customer_part_mappings m
JOIN customers c ON m.customer_id = c.id
JOIN parts p ON m.part_id = p.id;
""")
db_mappings = cur.fetchall()
print(f"Loaded {len(db_mappings)} customer_part_mappings")

# 4. Fetch DB customer_part_cross_references
cur.execute("""
SELECT c.customer_code, cr.customer_part_no, p.part_number, cr.part_id
FROM customer_part_cross_references cr
LEFT JOIN customers c ON cr.customer_id = c.id
LEFT JOIN parts p ON cr.part_id = p.id;
""")
db_cross_refs = cur.fetchall()
print(f"Loaded {len(db_cross_refs)} customer_part_cross_references")

# 5. Fetch DB po_master, orders, work_orders
cur.execute("SELECT id, po_number FROM po_master;")
db_pos = {r[1]: r[0] for r in cur.fetchall()}

cur.execute("SELECT id, oar_number FROM orders;")
db_oars = {r[1]: r[0] for r in cur.fetchall()}

cur.execute("SELECT id, wo_number FROM work_orders;")
db_wos = {r[1]: r[0] for r in cur.fetchall()}

print(f"\nDB State: {len(db_customers)} customers, {len(db_parts)} parts")
print(f"DB Operations: {len(db_pos)} POs, {len(db_oars)} OARs, {len(db_wos)} WOs")

print("\n" + "="*50)
print("ANALYZING WORKBOOK & DATASET")
print("="*50)

po_groups = defaultdict(list)
oar_groups = defaultdict(list)
wo_records = {}
duplicate_wos = []

for idx, r in enumerate(rows, 1):
    wo = str(r[0]).strip() if r[0] else None
    ts = str(r[1]) if r[1] else None
    oar = str(r[2]).strip() if r[2] else None
    cust = str(r[3]).strip() if r[3] else None
    cust_code = str(r[4]).strip() if r[4] else None
    uniq_prod_id = str(r[5]).strip() if r[5] else None
    part_no = str(r[6]).strip() if r[6] else None
    internal_part_code = str(r[7]).strip() if r[7] else None
    internal_id = str(r[8]).strip() if r[8] else None
    region = str(r[9]).strip() if r[9] else None
    grade = str(r[10]).strip() if r[10] else None
    cust_grade = str(r[11]).strip() if r[11] else None
    order_type = str(r[12]).strip() if r[12] else None
    order_cat = str(r[13]).strip() if r[13] else None
    po_num = str(r[14]).strip() if r[14] else None
    del_date = str(r[15]).strip() if r[15] else None
    po_rate = r[16]
    po_qty = r[17]
    wo_qty = r[18]
    remarks = str(r[19]).strip() if r[19] else ""
    wo_type = str(r[20]).strip() if r[20] else ""
    status = str(r[21]).strip() if r[21] else ""

    rec = {
        "row_idx": idx,
        "wo": wo,
        "ts": ts,
        "oar": oar,
        "cust": cust,
        "cust_code": cust_code,
        "uniq_prod_id": uniq_prod_id,
        "part_no": part_no,
        "internal_part_code": internal_part_code,
        "internal_id": internal_id,
        "region": region,
        "grade": grade,
        "cust_grade": cust_grade,
        "order_type": order_type,
        "order_cat": order_cat,
        "po_num": po_num,
        "del_date": del_date,
        "po_rate": po_rate,
        "po_qty": po_qty,
        "wo_qty": wo_qty,
        "remarks": remarks,
        "wo_type": wo_type,
        "status": status
    }
    
    po_groups[po_num].append(rec)
    oar_groups[oar].append(rec)
    if wo in wo_records:
        duplicate_wos.append(rec)
    else:
        wo_records[wo] = rec

print(f"\nTotal Source Rows: {len(rows)}")
print(f"Total Unique POs: {len(po_groups)}")
print(f"Total Unique OARs: {len(oar_groups)}")
print(f"Total Unique WOs: {len(wo_records)}")
print(f"Duplicate WOs: {len(duplicate_wos)}")

# Detailed PO breakdown
print("\n" + "="*50)
print("PO SUMMARY")
print("="*50)
for po, recs in po_groups.items():
    oars = sorted(list(set(r['oar'] for r in recs)))
    wos = [r['wo'] for r in recs]
    custs = set(r['cust'] for r in recs)
    cust_codes = set(r['cust_code'] for r in recs)
    del_dates = sorted(list(set(r['del_date'] for r in recs)))
    statuses = sorted(list(set(r['status'] for r in recs)))
    tot_wo_qty = sum(float(r['wo_qty'] or 0) for r in recs)
    po_qtys = set(r['po_qty'] for r in recs)
    print(f"\nPO Number: '{po}'")
    print(f"  Customer: {list(custs)} | Codes: {list(cust_codes)}")
    print(f"  OARs ({len(oars)}): {oars}")
    print(f"  WOs ({len(wos)}): {wos}")
    print(f"  PO Qty values in rows: {po_qtys} | Total WO Qty: {tot_wo_qty}")
    print(f"  Delivery Dates: {del_dates} | Statuses: {statuses}")

# Detailed OAR breakdown & Cross-PO check
print("\n" + "="*50)
print("OAR SUMMARY & CROSS-PO CHECK")
print("="*50)
for oar, recs in oar_groups.items():
    pos = sorted(list(set(r['po_num'] for r in recs)))
    cust_codes = set(r['cust_code'] for r in recs)
    wos = [r['wo'] for r in recs]
    tot_wo_qty = sum(float(r['wo_qty'] or 0) for r in recs)
    is_cross = len(pos) > 1
    print(f"OAR: '{oar}' | POs ({len(pos)}): {pos} | Cust: {list(cust_codes)} | WOs ({len(wos)}): {wos} | Tot WO Qty: {tot_wo_qty} | CROSS_PO_OAR: {is_cross}")

# Part Matching Check against DB
print("\n" + "="*50)
print("PART MATCHING CHECK")
print("="*50)
matched = []
unmatched = []

for r in rows:
    wo = str(r[0]).strip()
    cust_code = str(r[4]).strip() if r[4] else ""
    part_no = str(r[6]).strip() if r[6] else ""
    internal_part_code = str(r[7]).strip() if r[7] else ""
    internal_id = str(r[8]).strip() if r[8] else ""
    
    # Check parts table by internal_id
    p_match = None
    if internal_id in db_parts:
        p_match = ("MATCH_INTERNAL_ID", internal_id, db_parts[internal_id][0], db_parts[internal_id][1])
    elif part_no in db_parts:
        p_match = ("MATCH_PART_NO", part_no, db_parts[part_no][0], db_parts[part_no][1])
    else:
        # Check customer_part_mappings
        cur.execute("""
        SELECT p.id, p.part_number, m.customer_part_number
        FROM customer_part_mappings m
        JOIN parts p ON m.part_id = p.id
        JOIN customers c ON m.customer_id = c.id
        WHERE c.customer_code = %s AND (m.customer_part_number ILIKE %s OR p.part_number ILIKE %s);
        """, (cust_code, f"%{part_no}%", f"%{internal_id}%"))
        res = cur.fetchall()
        if res:
            p_match = ("MATCH_CUSTOMER_MAPPING", res)
            
    if p_match:
        matched.append((wo, cust_code, part_no, internal_id, internal_part_code, p_match))
    else:
        unmatched.append((wo, cust_code, part_no, internal_id, internal_part_code))

print(f"Total Rows: {len(rows)} | Matched: {len(matched)} | Unmatched: {len(unmatched)}")
if unmatched:
    print("\nUnmatched parts details:")
    for u in unmatched:
        print(f"  WO: {u[0]} | Cust: {u[1]} | Part No: '{u[2]}' | Internal ID: '{u[3]}' | Internal Part Code: '{u[4]}'")

cur.close()
conn.close()
connector.close()
