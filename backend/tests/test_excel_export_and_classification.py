import io
import pytest
import openpyxl
from fastapi.testclient import TestClient

from app.core.database import SessionLocal, auto_migrate_schema
from app.main import app
from app.services.excel_export_service import build_open_orders_and_wos_workbook, generate_excel_bytes
from app.models.order import Order, Customer, Part
from app.models.work_order import WorkOrder, WOStatus
from app.schemas.operations import OrderIntakeCreate, WorkOrderCreatePayload
from app.services.operations_service import OperationsService


@pytest.fixture(scope="module", autouse=True)
def setup_db():
    auto_migrate_schema()
    yield


def test_excel_export_workbook_structure():
    db = SessionLocal()
    try:
        wb = build_open_orders_and_wos_workbook(db)
        expected_sheets = [
            "Open_OARs_Dump",
            "Open_Work_Orders_Dump",
            "Traceability_Matrix",
            "Data_Capture_Specification",
            "Reference_Masters"
        ]
        assert wb.sheetnames == expected_sheets
        
        # Verify Sheet 1 headers
        ws_oars = wb["Open_OARs_Dump"]
        assert ws_oars.cell(row=4, column=1).value == "OAR Number"
        assert ws_oars.cell(row=4, column=2).value == "Classification"

        # Verify Sheet 2 headers
        ws_wos = wb["Open_Work_Orders_Dump"]
        assert ws_wos.cell(row=4, column=1).value == "WO Number"
        assert ws_wos.cell(row=4, column=3).value == "Classification"

        # Verify Data Capture Specification sheet has rows
        ws_spec = wb["Data_Capture_Specification"]
        assert ws_spec.max_row >= 15
        
        # Verify Reference Masters sheet has classifications
        ws_ref = wb["Reference_Masters"]
        assert ws_ref["A4"].value == "Production"
        assert ws_ref["A5"].value == "NPD"
    finally:
        db.close()


def test_order_intake_classification_inheritance():
    db = SessionLocal()
    try:
        # 1. Create NPD Order Intake with immediate WOs
        intake_npd = OrderIntakeCreate(
            customer_code="CUST-VALVE",
            customer_name="Flowserve Controls Ltd",
            customer_po="PO-NPD-TEST-001",
            part_number="BRZ-BUSH-100",
            po_quantity=100,
            order_type="NPD",
            create_wos=True,
            max_batch_size=50
        )
        res_npd = OperationsService.create_order_intake(db, intake_npd)
        assert res_npd.order_type == "NPD"
        assert len(res_npd.wos_created) == 2

        # Check in DB that order and child WOs have classification = 'NPD'
        db_order = db.query(Order).filter(Order.oar_number == res_npd.oar_number).first()
        assert db_order is not None
        assert db_order.order_type == "NPD"

        db_wos = db.query(WorkOrder).filter(WorkOrder.order_id == db_order.id).all()
        assert len(db_wos) == 2
        for w in db_wos:
            assert w.classification == "NPD"

        # 2. Create Production Order Intake
        intake_prod = OrderIntakeCreate(
            customer_code="CUST-PUMP",
            customer_name="Kirloskar Brothers Pumps",
            customer_po="PO-PROD-TEST-002",
            part_number="BRZ-RING-250",
            po_quantity=200,
            order_type="Production",
            create_wos=True,
            max_batch_size=100
        )
        res_prod = OperationsService.create_order_intake(db, intake_prod)
        assert res_prod.order_type == "Production"

        db_order_p = db.query(Order).filter(Order.oar_number == res_prod.oar_number).first()
        assert db_order_p.order_type == "Production"
        for w in db.query(WorkOrder).filter(WorkOrder.order_id == db_order_p.id).all():
            assert w.classification == "Production"
    finally:
        db.close()


def test_excel_export_bytes_and_filter():
    db = SessionLocal()
    try:
        # All export
        data_all = generate_excel_bytes(db)
        assert len(data_all) > 1000
        wb_all = openpyxl.load_workbook(io.BytesIO(data_all))
        assert "Open_OARs_Dump" in wb_all.sheetnames

        # NPD-only export
        data_npd = generate_excel_bytes(db, classification_filter="NPD")
        assert len(data_npd) > 1000
        wb_npd = openpyxl.load_workbook(io.BytesIO(data_npd))
        ws = wb_npd["Open_OARs_Dump"]
        for row in range(5, ws.max_row + 1):
            val = ws.cell(row=row, column=2).value
            if val:
                assert val.upper() == "NPD"
    finally:
        db.close()
