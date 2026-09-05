"""
VSPL — Stage-wise WIP Reconciliation Checker (v2, for oms_engine v3 reports)
=============================================================================
Cross-checks three independent views of WIP in a Daily_Report_<date>.xlsx:

  3a Inventory by Stage   — stage totals (v3: built bottom-up from per-WO figures)
  3b Inventory WO+Part    — per-WO "WIP Detail" / "WIP Total"
  3c Stage Audit Sheet    — one row per WO per stage holding WIP (the floor-audit sheet)

In v3 all three come from the same route-aware per-WO calculation
    WIP after S = OK(S) - (OK(next on route) + Rej(next on route))
so they MUST agree; any mismatch means the report file was edited by hand or is stale.
The checker also accepts a completed physical audit: if the auditor filled the
"Physical Count" column on 3c, book-vs-floor variances are computed and flagged.

v2 changes vs v1:
  * stage list is read DYNAMICALLY from the report (v3 stages are extensible via the
    WO Release input — never hardcode F1..FI)
  * reconciles 3a vs 3c (and 3b totals), not the old clamped plant-wide rollup
  * writes its findings to a new "WIP Reconciliation" sheet in a copy of the workbook

Run
---
    python wip_reconcile.py Daily_Report_2026-07-14.xlsx [--out RECONCILED.xlsx]
"""

import argparse
import os
import sys
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BLUE = "2E5496"
GREEN_FILL = PatternFill("solid", fgColor="C6E0B4")
RED_FILL = PatternFill("solid", fgColor="F4C7C3")
AMBER_FILL = PatternFill("solid", fgColor="FFE699")
HF = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HFILL = PatternFill("solid", fgColor=BLUE)


def reconcile(df_3a, df_3b, df_3c):
    stages = [str(s) for s in df_3a["After Stage"].dropna()]  # dynamic, report-ordered
    a = {str(r["After Stage"]): int(pd.to_numeric(r.get("WIP On-hand"), errors="coerce") or 0)
         for _, r in df_3a.iterrows()}
    c = (df_3c.groupby("Stage")["WIP Qty (Book)"].sum().astype(int).to_dict()
         if len(df_3c) else {})
    b_total = int(pd.to_numeric(df_3b.get("WIP Total"), errors="coerce").fillna(0).sum()) \
        if "WIP Total" in df_3b.columns else None

    rows = []
    for s in stages:
        ra, rc = a.get(s, 0), int(c.get(s, 0))
        diff = rc - ra
        rows.append({"Stage": s, "3a Stage Total": ra, "3c Audit-Sheet Sum": rc,
                     "Difference (3c-3a)": diff,
                     "Status": "MATCH" if diff == 0 else "MISMATCH",
                     "Note": "" if diff == 0 else
                     "3a and 3c disagree — the report was hand-edited or is stale; regenerate it."})
    total_a = sum(a.get(s, 0) for s in stages)
    rows.append({"Stage": "TOTAL", "3a Stage Total": total_a,
                 "3c Audit-Sheet Sum": sum(int(v) for v in c.values()),
                 "Difference (3c-3a)": sum(int(v) for v in c.values()) - total_a,
                 "Status": "MATCH" if (b_total is None or b_total == total_a) else "MISMATCH",
                 "Note": "" if (b_total is None or b_total == total_a) else
                 f"3b per-WO totals sum to {b_total}, 3a to {total_a}."})
    return pd.DataFrame(rows)


def physical_variances(df_3c):
    """If the auditor filled 'Physical Count' on 3c, compute book-vs-floor variances."""
    if not len(df_3c) or "Physical Count" not in df_3c.columns:
        return pd.DataFrame()
    d = df_3c.copy()
    d["Physical Count"] = pd.to_numeric(d["Physical Count"], errors="coerce")
    d = d[d["Physical Count"].notna()]
    if not len(d):
        return pd.DataFrame()
    d["Book"] = pd.to_numeric(d["WIP Qty (Book)"], errors="coerce").fillna(0).astype(int)
    d["Variance (Phys-Book)"] = d["Physical Count"].astype(int) - d["Book"]
    return d[["Stage", "WO ID", "Part No", "Book", "Physical Count", "Variance (Phys-Book)"]]


def _style(ws, status_col, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HF
        cell.fill = HFILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = 20
    ws.column_dimensions[get_column_letter(ncols)].width = 70
    for r in range(2, ws.max_row + 1):
        v = str(ws.cell(row=r, column=status_col).value)
        fill = GREEN_FILL if v == "MATCH" else (RED_FILL if v == "MISMATCH" else AMBER_FILL)
        for c in range(1, ncols + 1):
            ws.cell(row=r, column=c).fill = fill
    ws.freeze_panes = "A2"


def main():
    ap = argparse.ArgumentParser(description="VSPL stage-wise WIP reconciliation checker v2")
    ap.add_argument("report", help="Path to Daily_Report_<date>.xlsx (oms_engine v3)")
    ap.add_argument("--out", help="Output path (default: <report>_RECONCILED.xlsx)")
    args = ap.parse_args()
    if not os.path.exists(args.report):
        sys.exit(f"File not found: {args.report}")
    out_path = args.out or (os.path.splitext(args.report)[0] + "_RECONCILED.xlsx")

    df_3a = pd.read_excel(args.report, sheet_name="3a Inventory by Stage")
    df_3b = pd.read_excel(args.report, sheet_name="3b Inventory WO+Part")
    try:
        df_3c = pd.read_excel(args.report, sheet_name="3c Stage Audit Sheet")
    except Exception:
        sys.exit("No '3c Stage Audit Sheet' — this report predates oms_engine v3. "
                 "Re-run the daily cycle with the v3 engine first.")

    recon = reconcile(df_3a, df_3b, df_3c)
    phys = physical_variances(df_3c)

    wb = load_workbook(args.report)
    for name in ("WIP Reconciliation", "Physical Variances"):
        if name in wb.sheetnames:
            del wb[name]
    ws = wb.create_sheet("WIP Reconciliation")
    ws.append(list(recon.columns))
    for _, row in recon.iterrows():
        ws.append(list(row))
    _style(ws, status_col=5, ncols=len(recon.columns))
    if len(phys):
        wsp = wb.create_sheet("Physical Variances")
        wsp.append(list(phys.columns))
        for _, row in phys.iterrows():
            wsp.append(list(row))
        for c in range(1, len(phys.columns) + 1):
            wsp.cell(row=1, column=c).font = HF
            wsp.cell(row=1, column=c).fill = HFILL
            wsp.column_dimensions[get_column_letter(c)].width = 16
        for r in range(2, wsp.max_row + 1):
            var = wsp.cell(row=r, column=len(phys.columns)).value
            fill = GREEN_FILL if var == 0 else RED_FILL
            for c in range(1, len(phys.columns) + 1):
                wsp.cell(row=r, column=c).fill = fill
        wsp.freeze_panes = "A2"
    wb.save(out_path)

    print("Stage-wise WIP Reconciliation (3a vs 3c bottom-up)")
    print("=" * 64)
    for _, r in recon.iterrows():
        flag = "OK" if r["Status"] == "MATCH" else "!! MISMATCH"
        print(f"{r['Stage']:>6}  3a={r['3a Stage Total']:>6}  3c={r['3c Audit-Sheet Sum']:>6}"
              f"  diff={r['Difference (3c-3a)']:>5}  {flag}")
    n_mm = (recon["Status"] == "MISMATCH").sum()
    if len(phys):
        n_var = int((phys["Variance (Phys-Book)"] != 0).sum())
        print(f"Physical audit: {len(phys)} counted line(s), {n_var} variance(s) — "
              f"see 'Physical Variances'.")
    print("=" * 64)
    print(("All stages reconcile cleanly." if n_mm == 0 else f"{n_mm} stage(s) need review.")
          + f" Output: {out_path}")


if __name__ == "__main__":
    main()
